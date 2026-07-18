from __future__ import annotations

import base64
import io
import zipfile

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.extract.files import MAX_ATTACHMENT_BYTES, UnsupportedFile, extract


def b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode("ascii")


def xlsx_bytes_with_shared_strings(*values: str) -> bytes:
    shared_items = "".join(f'<si><t>{value}</t></si>' for value in values)
    cells = "".join(f'<c r="A{idx + 1}" t="s"><v>{idx}</v></c>' for idx, _ in enumerate(values))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/workbook.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>
</workbook>""",
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1">{cells}</row></sheetData>
</worksheet>""",
        )
        archive.writestr(
            "xl/sharedStrings.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(values)}" uniqueCount="{len(values)}">{shared_items}</sst>""",
        )
    return buf.getvalue()


def test_extract_text_plain() -> None:
    doc = extract("note.txt", "text/plain", b64(b"My IC is 880412-08-5561"))
    assert doc.text == "My IC is 880412-08-5561"
    assert doc.source_filename == "note.txt"
    assert doc.mime == "text/plain"


def test_extract_csv_as_utf8_text() -> None:
    doc = extract("payroll.csv", "text/csv", b64(b"name,salary\nAli,RM 8,500"))
    assert "RM 8,500" in doc.text


def test_extract_docx_paragraphs_and_tables() -> None:
    source = DocxDocument()
    source.add_paragraph("Employee IC 880412-08-5561")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "salary"
    table.cell(0, 1).text = "RM 8,500"
    buf = io.BytesIO()
    source.save(buf)

    doc = extract(
        "employee.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        b64(buf.getvalue()),
    )

    assert "Employee IC 880412-08-5561" in doc.text
    assert "salary\tRM 8,500" in doc.text


def test_extract_docx_headers_and_footers() -> None:
    source = DocxDocument()
    source.add_paragraph("Body text")
    source.sections[0].header.paragraphs[0].text = "Header IC 880412-08-5561"
    source.sections[0].footer.paragraphs[0].text = "Footer salary RM 8,500"
    buf = io.BytesIO()
    source.save(buf)

    doc = extract(
        "employee.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        b64(buf.getvalue()),
    )

    assert "Header IC 880412-08-5561" in doc.text
    assert "Footer salary RM 8,500" in doc.text


def test_extract_pptx_text_shapes() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quarterly revenue is below forecast."
    buf = io.BytesIO()
    deck.save(buf)

    doc = extract(
        "finance.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        b64(buf.getvalue()),
    )

    assert "Quarterly revenue is below forecast." in doc.text


def test_extract_pptx_speaker_notes() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Public title"
    slide.notes_slide.notes_text_frame.text = "Speaker IC 880412-08-5561"
    buf = io.BytesIO()
    deck.save(buf)

    doc = extract(
        "speaker-notes.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        b64(buf.getvalue()),
    )

    assert "Public title" in doc.text
    assert "Speaker IC 880412-08-5561" in doc.text


def test_extract_xlsx_shared_strings() -> None:
    doc = extract(
        "payroll.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        b64(xlsx_bytes_with_shared_strings("Employee IC 880412-08-5561", "RM 8,500")),
    )

    assert "Employee IC 880412-08-5561" in doc.text
    assert "RM 8,500" in doc.text


def test_extract_pdf_rejects_invalid_pdf() -> None:
    with pytest.raises(UnsupportedFile):
        extract("broken.pdf", "application/pdf", b64(b"not a pdf"))


def test_unsupported_file_type() -> None:
    with pytest.raises(UnsupportedFile):
        extract("image.png", "image/png", b64(b"whatever"))


def test_invalid_base64() -> None:
    with pytest.raises(UnsupportedFile):
        extract("note.txt", "text/plain", "not-base64")


def test_size_limit() -> None:
    too_large = b64(b"x" * (MAX_ATTACHMENT_BYTES + 1))
    with pytest.raises(UnsupportedFile):
        extract("large.txt", "text/plain", too_large)
