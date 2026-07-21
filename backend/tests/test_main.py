"""M1 pipeline tests: one per verdict, the hold->approve->unblock loop, the
error mappings, SD-1, and the ENV=demo guard. Tests run with the configured
local module frame and REPO_BACKEND=memory -- no Docker.
"""

from __future__ import annotations

import base64
import io
import zipfile
from datetime import datetime, timedelta, timezone
from xml.sax.saxutils import escape

import numpy as np
import pdfplumber
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.extract.files import extract
from app.main import API_KEY, _check_demo_guard, _pdf_bytes_from_text, app
from app.repo import reset_memory_repos

HEADERS = {"X-Api-Key": API_KEY}


def xlsx_bytes_with_shared_strings(*values: str) -> bytes:
    shared_items = "".join(f'<si><t>{value}</t></si>' for value in values)
    cells = "".join(f'<c r="A{idx + 1}" t="s"><v>{idx}</v></c>' for idx, _ in enumerate(values))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "xl/workbook.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>
</workbook>""",
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1">{cells}</row></sheetData>
</worksheet>""",
        )
        archive.writestr(
            "xl/sharedStrings.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(values)}" uniqueCount="{len(values)}">{shared_items}</sst>""",
        )
    return buf.getvalue()


def image_bytes_with_text(text: str) -> bytes:
    image = Image.new("RGB", (1200, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 52)
    draw.text((32, 72), text, fill="black", font=font)
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def scanned_pdf_bytes_with_text(text: str) -> bytes:
    image = Image.open(io.BytesIO(image_bytes_with_text(text)))
    buf = io.BytesIO()
    image.save(buf, format="PDF", resolution=150.0)
    return buf.getvalue()


def docx_bytes_with_comments_and_revisions(comment: str, inserted: str, deleted: str) -> bytes:
    source = DocxDocument()
    source.add_paragraph("Body text")
    buf = io.BytesIO()
    source.save(buf)

    comment_reference_xml = '<w:p><w:r><w:commentReference w:id="0"/></w:r></w:p>'
    revision_xml = (
        comment_reference_xml
        +
        f'<w:p><w:ins w:id="1" w:author="test"><w:r><w:t>{escape(inserted)}</w:t></w:r></w:ins></w:p>'
        f'<w:p><w:del w:id="2" w:author="test"><w:r><w:delText>{escape(deleted)}</w:delText></w:r></w:del></w:p>'
    )
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        rels_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        content_types_xml = archive.read("[Content_Types].xml").decode("utf-8")
    if "<w:sectPr" in document_xml:
        document_xml = document_xml.replace("<w:sectPr", revision_xml + "<w:sectPr", 1)
    else:
        document_xml = document_xml.replace("</w:body>", revision_xml + "</w:body>", 1)

    comments_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:comment w:id="0" w:author="test" w:date="2026-07-19T00:00:00Z">
<w:p><w:r><w:t>{escape(comment)}</w:t></w:r></w:p>
</w:comment>
</w:comments>"""
    if "relationships/comments" not in rels_xml:
        rels_xml = rels_xml.replace(
            "</Relationships>",
            '<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/></Relationships>',
            1,
        )
    if "word/comments.xml" not in content_types_xml:
        content_types_xml = content_types_xml.replace(
            "</Types>",
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>',
            1,
        )

    return replace_zip_parts(
        buf.getvalue(),
        {
            "[Content_Types].xml": content_types_xml.encode("utf-8"),
            "word/document.xml": document_xml.encode("utf-8"),
            "word/_rels/document.xml.rels": rels_xml.encode("utf-8"),
            "word/comments.xml": comments_xml.encode("utf-8"),
        },
    )


def pptx_bytes_with_comment(comment: str) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Public title"
    buf = io.BytesIO()
    deck.save(buf)

    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as archive:
        rels_xml = archive.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8")
        content_types_xml = archive.read("[Content_Types].xml").decode("utf-8")

    comments_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cm authorId="0" dt="2026-07-19T00:00:00Z" idx="1">
<p:text>{escape(comment)}</p:text>
</p:cm>
</p:cmLst>"""
    if "relationships/comments" not in rels_xml:
        rels_xml = rels_xml.replace(
            "</Relationships>",
            '<Relationship Id="rId999" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="../comments/comment1.xml"/></Relationships>',
            1,
        )
    if "ppt/comments/comment1.xml" not in content_types_xml:
        content_types_xml = content_types_xml.replace(
            "</Types>",
            '<Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/></Types>',
            1,
        )

    return replace_zip_parts(
        buf.getvalue(),
        {
            "[Content_Types].xml": content_types_xml.encode("utf-8"),
            "ppt/slides/_rels/slide1.xml.rels": rels_xml.encode("utf-8"),
            "ppt/comments/comment1.xml": comments_xml.encode("utf-8"),
        },
    )


def replace_zip_parts(raw: bytes, replacements: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest:
        written = set()
        for info in source.infolist():
            content = replacements.get(info.filename, source.read(info.filename))
            dest.writestr(info, content)
            written.add(info.filename)
        for filename, content in replacements.items():
            if filename not in written:
                dest.writestr(filename, content)
    return output.getvalue()


def docx_bytes_with_embedded_docx(text: str) -> bytes:
    embedded = DocxDocument()
    embedded.add_paragraph(text)
    embedded_buf = io.BytesIO()
    embedded.save(embedded_buf)

    outer = DocxDocument()
    outer.add_paragraph("Outer document")
    outer_buf = io.BytesIO()
    outer.save(outer_buf)
    return replace_zip_parts(
        outer_buf.getvalue(),
        {"word/embeddings/embedded.docx": embedded_buf.getvalue()},
    )


def docx_bytes_with_legacy_ole_object() -> bytes:
    source = DocxDocument()
    source.add_paragraph("Outer document")
    buf = io.BytesIO()
    source.save(buf)
    return replace_zip_parts(
        buf.getvalue(),
        {"word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0legacy ole payload"},
    )


def pptx_bytes_with_embedded_docx(text: str) -> bytes:
    embedded = DocxDocument()
    embedded.add_paragraph(text)
    embedded_buf = io.BytesIO()
    embedded.save(embedded_buf)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Outer deck"
    deck_buf = io.BytesIO()
    deck.save(deck_buf)
    return replace_zip_parts(
        deck_buf.getvalue(),
        {"ppt/embeddings/embedded.docx": embedded_buf.getvalue()},
    )


def pptx_bytes_with_legacy_ole_object() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Outer deck"
    buf = io.BytesIO()
    deck.save(buf)
    return replace_zip_parts(
        buf.getvalue(),
        {"ppt/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0legacy ole payload"},
    )


@pytest.fixture(autouse=True)
def _fresh_repos():
    """The memory backend is a process-wide singleton (state has to survive
    between requests) -- reset it around every test so they cannot see each
    other's audit rows."""
    reset_memory_repos()
    yield
    reset_memory_repos()


@pytest.fixture
def client():
    return TestClient(app)


def inspect(client: TestClient, **overrides):
    body = {
        "user_id": "wei.eng@acme.my",
        "target_host": "chatgpt.com",
        "prompt": "hello",
        **overrides,
    }
    return client.post("/v1/inspect", headers=HEADERS, json=body)


def seed_internal_doc_for_m8(
    client: TestClient, monkeypatch: pytest.MonkeyPatch, *,
    department: str = "hr", similarity: float = 0.88,
) -> None:
    """M10: EmbedDetector._best is stubbed to a fixed (department, score)
    pair so tests can land precisely inside a policy tier band -- real
    cosine similarity from two identical constant vectors is always exactly
    1.0 and cannot express "0.88, in the hold band but below deny"."""
    class FakeEncoder:
        def encode(self, texts: list[str], *, normalize_embeddings: bool) -> np.ndarray:
            assert normalize_embeddings is True
            return np.asarray([[1.0] + [0.0] * 383 for _ in texts])

    monkeypatch.setattr("app.main.get_m8_encoder", lambda: FakeEncoder())
    monkeypatch.setattr(
        "app.detect.embed.EmbedDetector._load_default_encoder",
        staticmethod(lambda: FakeEncoder()),
    )
    monkeypatch.setattr(
        "app.detect.embed.EmbedDetector._best",
        lambda self, text: (department, similarity),
    )

    content = base64.b64encode(
        b"HR confidential termination plan for employee grievance cases."
    ).decode("ascii")
    r = client.post(
        "/v1/internal-docs",
        headers=HEADERS,
        json={
            "department": department,
            "attachment": {
                "filename": "hr-termination-plan.txt",
                "mime": "text/plain",
                "content_b64": content,
            },
        },
    )
    assert r.status_code == 200


def mock_deberta_non_confidential(monkeypatch: pytest.MonkeyPatch) -> None:
    """Neutralise M9 for tests that need to pin the verdict to M8's mocked
    confidence alone. This repo's .env sets SEMANTIC_DETECTOR=both, so the
    REAL DeBERTa model runs on every request in this test file too -- and
    since CompositeSemanticDetector.detect() dedups overlapping same-class
    spans by keeping only the higher-confidence one, a real DeBERTa span can
    silently replace the mocked M8 span before decide() ever sees it,
    defeating the whole point of mocking M8. This stub short-circuits model
    loading entirely (no weights read) and always reports "not confidential",
    so M8's mocked span is the only DEPT_CONFIDENTIAL signal in play. Scoped
    (via monkeypatch, auto-reverted) to only the tests that call this --
    every other test still exercises the real DeBERTa model."""
    from app.detect.deberta import SentencePrediction

    monkeypatch.setattr("app.detect.deberta.DebertaDetector._load_model", lambda self: None)
    monkeypatch.setattr(
        "app.detect.deberta.DebertaDetector._predict",
        lambda self, sentence: SentencePrediction("other", 0.0, False, 0.0),
    )


def seed_internal_doc_for_m8_hold(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    """Backwards-compatible name for the hold/approve/deny/timeout review
    workflow tests below, which only need SOME hold verdict and do not care
    about the exact similarity band. Mocks DeBERTa too (see
    mock_deberta_non_confidential) so that hold verdict is reliably M8's."""
    mock_deberta_non_confidential(monkeypatch)
    seed_internal_doc_for_m8(client, monkeypatch, similarity=0.88)


# --------------------------------------------------------------------------
# One test per verdict
# --------------------------------------------------------------------------


def test_allow(client: TestClient) -> None:
    r = inspect(client, prompt="Turn this into bullet points for the standup.")
    assert r.status_code == 200
    body = r.json()
    assert body["decision"] == "allow"
    assert body["redacted_prompt"] is None
    assert body["redacted_attachments"] == []
    assert body["data_classes"] == []
    assert body["content_department"] is None
    assert body["department_confidence"] == 0.0
    assert body["debug"] == {
        "policy_layer": "M3",
        "semantic_detector": "m8",
        "detections": [],
    }

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    assert any(x["log_id"] == body["log_id"] for x in logs["items"])


def test_secret_is_redacted_not_denied(client: TestClient) -> None:
    # M2 has no policy voice (see rules.yaml): a leaked credential is
    # redacted by the structural regex floor exactly like any other regex
    # class, never escalated to deny. Only DEPT_CONFIDENTIAL (M8/M9) can
    # push a request to hold/deny now.
    r = inspect(client,
                prompt="Debug this: OPENAI_API_KEY=sk-proj-9f2b1c8e4a7d3f6b0e5c2a9d")
    body = r.json()
    assert body["decision"] == "redact"
    assert "SECRET" in body["data_classes"]
    assert body["redacted_prompt"] is not None
    assert "sk-proj-9f2b1c8e4a7d3f6b0e5c2a9d" not in body["redacted_prompt"]


def test_redact(client: TestClient) -> None:
    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt="Summarise this employee grievance from Encik Rahman, "
               "IC 880412-08-5561.",
    )
    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert body["redacted_prompt"] is not None
    assert "880412-08-5561" not in body["redacted_prompt"]
    assert "[MYKAD_1]" in body["redacted_prompt"]


def test_salary_is_redacted_no_purpose_gate(client: TestClient) -> None:
    # M2 has no policy row any more: no require_purpose gate. It still
    # always gets redacted via the structural regex floor.
    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt="Draft a reply to the payroll query. "
               "Gaji bulanan RM 8,500 for grade M3.",
    )
    body = r.json()
    assert body["decision"] == "redact"
    assert body["require_purpose"] is False
    assert "SALARY" in body["data_classes"]
    assert body["debug"]["detections"] == [
        {
            "data_class": "SALARY",
            "detector": "rules",
            "confidence": 1.0,
            "source": "prompt",
        }
    ]
    assert body["redacted_prompt"] is not None
    assert "RM 8,500" not in body["redacted_prompt"]
    assert "[SALARY_1]" in body["redacted_prompt"]


def test_m8_internal_doc_mid_high_similarity_holds_with_purpose(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # M10: DEPT_CONFIDENTIAL never redacts on PUBLIC_AI any more -- M8's
    # tiers are allow+justify / hold+justify / deny. 0.88 similarity lands
    # in the hold band regardless of department (same department here).
    seed_internal_doc_for_m8(client, monkeypatch, department="hr", similarity=0.88)

    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt="Please summarize the HR confidential termination plan "
               "for employee grievance cases.",
    )
    body = r.json()
    assert body["decision"] == "hold"
    assert body["require_purpose"] is True
    assert "DEPT_CONFIDENTIAL" in body["data_classes"]
    assert body["redacted_prompt"] is None


def test_m8_internal_doc_cross_department_holds_from_the_same_tier_row(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # M10 deleted the @cross policy row -- cross-department mismatch (an
    # engineer reading HR-indexed content) is display/log only now. This
    # lands on the exact same hold band, from the same tier row, as the
    # same-department case above; the department mismatch changes nothing
    # about the decision.
    mock_deberta_non_confidential(monkeypatch)
    seed_internal_doc_for_m8(client, monkeypatch, department="hr", similarity=0.88)

    r = inspect(
        client, user_id="wei.eng@acme.my",
        prompt="Please summarize the HR confidential termination plan "
               "for employee grievance cases.",
    )
    body = r.json()
    assert body["decision"] == "hold"
    assert body["require_purpose"] is True
    assert "DEPT_CONFIDENTIAL" in body["data_classes"]
    assert body["redacted_prompt"] is None
    assert body["content_department"] == "hr"


def test_m8_internal_doc_very_high_similarity_denies(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # M10: near-verbatim internal-document exfiltration (M8 similarity
    # >= 0.92) is a deny -- redaction cannot un-leak an exact-match document.
    seed_internal_doc_for_m8(client, monkeypatch, department="hr", similarity=0.94)

    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt="Please summarize the HR confidential termination plan "
               "for employee grievance cases.",
    )
    body = r.json()
    assert body["decision"] == "deny"
    assert body["require_purpose"] is False
    assert "DEPT_CONFIDENTIAL" in body["data_classes"]
    assert body["redacted_prompt"] is None


def test_m8_high_similarity_in_attachment_only_still_denies(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Regression: a DEPT_CONFIDENTIAL span found ONLY in an attachment (no
    # echo in the prompt text) must still reach decide()'s per-detector tier
    # table. Classification.spans used to be prompt-only, so an attachment-
    # only signal was invisible to embed/deberta-scoped rules (which have no
    # classification.confidence fallback -- that applies only to "*" rules)
    # and silently defaulted to allow.
    seed_internal_doc_for_m8(client, monkeypatch, department="hr", similarity=0.94)

    content = base64.b64encode(
        b"HR confidential termination plan for employee grievance cases."
    ).decode("ascii")
    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt="see attached, nothing sensitive typed here",
        attachments=[{
            "filename": "plan.txt", "mime": "text/plain", "content_b64": content,
        }],
    )
    body = r.json()
    assert body["decision"] == "deny"
    assert "DEPT_CONFIDENTIAL" in body["data_classes"]


# --------------------------------------------------------------------------
# SD-1: the original value must never reach the audit row.
# --------------------------------------------------------------------------


def test_sd1_original_never_in_audit_record(client: TestClient) -> None:
    original = "880412-08-5561"
    r = inspect(
        client, user_id="aiman.hr@acme.my",
        prompt=f"Summarise this employee grievance from Encik Rahman, IC {original}.",
    )
    log_id = r.json()["log_id"]

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    row = next(x for x in logs["items"] if x["log_id"] == log_id)

    assert original not in row["prompt_text"]
    assert row["map_id"] is not None


# --------------------------------------------------------------------------
# THE loop: hold -> approve -> unblock.
# --------------------------------------------------------------------------


def test_hold_approve_unblock_loop(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    seed_internal_doc_for_m8_hold(client, monkeypatch)

    prompt_text = ("Please summarize the HR confidential termination plan "
                   "for employee grievance cases.")
    r = inspect(client, user_id="wei.eng@acme.my", prompt=prompt_text)
    body = r.json()
    log_id = body["log_id"]
    assert body["decision"] == "hold"

    status = client.get(f"/v1/inspect/{log_id}", headers=HEADERS).json()
    assert status == {"decision": "hold", "status": "pending"}

    pending = client.get("/v1/logs", headers=HEADERS,
                         params={"status": "pending", "limit": 200}).json()
    assert any(x["log_id"] == log_id for x in pending["items"])

    dec = client.post(f"/v1/logs/{log_id}/decision", headers=HEADERS,
                      json={"approve": True, "remark": "aggregate figures only"})
    assert dec.status_code == 200
    assert dec.json()["status"] == "approved"
    assert dec.json()["decision"] == "hold"   # the verdict is a historical fact

    # A human-reviewed approval now returns the prepared redaction too, same
    # as the lazy timeout downgrade -- this prompt has no M2 span, so nothing
    # needed masking and redacted_prompt equals the original verbatim.
    unblocked = client.get(f"/v1/inspect/{log_id}", headers=HEADERS).json()
    assert unblocked == {"decision": "hold", "status": "approved", "redacted_prompt": prompt_text}


def test_hold_deny_review_sets_status_denied(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    seed_internal_doc_for_m8_hold(client, monkeypatch)

    r = inspect(
        client, user_id="wei.eng@acme.my",
        prompt="Please summarize the HR confidential termination plan "
               "for employee grievance cases.",
    )
    assert r.json()["decision"] == "hold"
    log_id = r.json()["log_id"]

    dec = client.post(f"/v1/logs/{log_id}/decision", headers=HEADERS,
                      json={"approve": False, "remark": "use the internal assistant"})
    assert dec.json()["status"] == "denied"

    status = client.get(f"/v1/inspect/{log_id}", headers=HEADERS).json()
    assert status["status"] == "denied"


# --------------------------------------------------------------------------
# Shadow AI: SD-7, auto-registered, never held on that account alone.
# --------------------------------------------------------------------------


def test_shadow_ai_registers_and_does_not_hold(client: TestClient) -> None:
    r = inspect(client, target_host="deepseek.com",
               prompt="harmless prompt about nothing sensitive")
    assert r.json()["decision"] == "allow"

    tools = client.get("/v1/tools", headers=HEADERS).json()
    row = next(t for t in tools if t["host"] == "deepseek.com")
    assert row["status"] == "unapproved"
    assert row["tier"] == "PUBLIC_AI"
    assert row["usage_count"] >= 1


def test_tools_include_first_seen(client: TestClient) -> None:
    tools = client.get("/v1/tools", headers=HEADERS).json()
    seeded = next(t for t in tools if t["host"] == "chatgpt.com")
    assert seeded["first_seen"] is not None  # rides along on asdict() for free

    inspect(client, target_host="brand-new-shadow.ai", prompt="hi")
    tools = client.get("/v1/tools", headers=HEADERS).json()
    fresh = next(t for t in tools if t["host"] == "brand-new-shadow.ai")
    assert fresh["first_seen"] is not None


def test_internal_doc_upload_indexes_extracted_chunks(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeEncoder:
        def encode(self, texts: list[str], *, normalize_embeddings: bool) -> np.ndarray:
            assert normalize_embeddings is True
            return np.asarray([[1.0] + [0.0] * 383 for _ in texts])

    monkeypatch.setattr("app.main.get_m8_encoder", lambda: FakeEncoder())

    content = base64.b64encode(
        b"2027 Product Roadmap is confidential engineering planning."
    ).decode("ascii")
    r = client.post(
        "/v1/internal-docs",
        headers=HEADERS,
        json={
            "department": "engineering",
            "attachment": {
                "filename": "roadmap.txt",
                "mime": "text/plain",
                "content_b64": content,
            },
        },
    )

    assert r.status_code == 200
    body = r.json()
    assert body["source_file"] == "roadmap.txt"
    assert body["department"] == "engineering"
    assert body["chunk_count"] == 1

    listed = client.get("/v1/internal-docs", headers=HEADERS).json()
    assert listed["total"] == 1
    assert listed["items"][0]["source_file"] == "roadmap.txt"
    assert "2027 Product Roadmap" in listed["items"][0]["chunk_text"]
    assert "embedding" not in listed["items"][0]


# --------------------------------------------------------------------------
# cross_department_only (M6's cross-department view)
# --------------------------------------------------------------------------


def test_logs_cross_department_only_filters_and_totals_correctly(client: TestClient) -> None:
    def is_cross_department(row: dict) -> bool:
        return (
            row["content_department"] is not None
            and row["user_department"] is not None
            and row["content_department"] != row["user_department"]
            and row["department_confidence"] > 0.6
        )

    everything = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    assert everything["total"] == 30    # the seeded set, unfiltered
    expected_ids = {r["log_id"] for r in everything["items"] if is_cross_department(r)}
    assert expected_ids   # the seed plants mismatches on purpose

    filtered = client.get("/v1/logs", headers=HEADERS,
                          params={"limit": 200, "cross_department_only": "true"}).json()
    assert filtered["total"] == len(expected_ids)
    assert {r["log_id"] for r in filtered["items"]} == expected_ids


# --------------------------------------------------------------------------
# Errors
# --------------------------------------------------------------------------


def test_unknown_user_is_403(client: TestClient) -> None:
    r = inspect(client, user_id="nobody@acme.my")
    assert r.status_code == 403


def test_unsupported_file_is_415(client: TestClient) -> None:
    r = inspect(client, prompt="see attached", attachments=[
        {"filename": "x.pdf", "mime": "application/pdf",
         "content_b64": base64.b64encode(b"whatever").decode()},
    ])
    assert r.status_code == 415


def test_docx_legacy_ole_embedded_object_is_415(client: TestClient) -> None:
    r = inspect(client, prompt="see attached", attachments=[
        {
            "filename": "outer.docx",
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "content_b64": base64.b64encode(docx_bytes_with_legacy_ole_object()).decode(),
        },
    ])
    assert r.status_code == 415


def test_pptx_legacy_ole_embedded_object_is_415(client: TestClient) -> None:
    r = inspect(client, prompt="see attached", attachments=[
        {
            "filename": "outer.pptx",
            "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "content_b64": base64.b64encode(pptx_bytes_with_legacy_ole_object()).decode(),
        },
    ])
    assert r.status_code == 415


def test_unknown_log_id_is_404(client: TestClient) -> None:
    r = client.get("/v1/inspect/does-not-exist", headers=HEADERS)
    assert r.status_code == 404

    r2 = client.post("/v1/logs/does-not-exist/decision", headers=HEADERS,
                     json={"approve": True, "remark": "x"})
    assert r2.status_code == 404

    r3 = client.post("/v1/tools/does-not-exist.example/decision", headers=HEADERS,
                     json={"status": "approved", "remark": "x"})
    assert r3.status_code == 404


def test_wrong_or_missing_api_key_is_403(client: TestClient) -> None:
    assert client.get("/v1/health", headers={"X-Api-Key": "wrong"}).status_code == 403
    assert client.get("/v1/health").status_code == 403


# --------------------------------------------------------------------------
# Text extraction round-trip through the mock plane.
# --------------------------------------------------------------------------


def test_supported_attachment_contributes_to_verdict(client: TestClient) -> None:
    content = base64.b64encode(b"RM 8,500 buried in a csv cell").decode()
    r = inspect(client, user_id="aiman.hr@acme.my",
               prompt="see attached", attachments=[
                   {"filename": "note.csv", "mime": "text/csv",
                    "content_b64": content},
               ])
    body = r.json()
    assert body["decision"] == "redact"
    assert body["require_purpose"] is False
    assert "SALARY" in body["data_classes"]
    assert body["debug"]["detections"] == [
        {
            "data_class": "SALARY",
            "detector": "rules",
            "confidence": 1.0,
            "source": "attachment",
        }
    ]
    assert len(body["redacted_attachments"]) == 1


def test_docx_attachment_contributes_to_verdict(client: TestClient) -> None:
    docx = DocxDocument()
    docx.add_paragraph("Compensation note: Gaji bulanan RM 8,500.")
    buf = io.BytesIO()
    docx.save(buf)

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "salary.docx",
                        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "content_b64": base64.b64encode(buf.getvalue()).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert body["require_purpose"] is False
    assert "SALARY" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1


def test_text_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    original = "Employee IC 880412-08-5561 needs a reply."
    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "employee.txt",
                        "mime": "text/plain",
                        "content_b64": base64.b64encode(original.encode()).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert body["redacted_prompt"] == "see attached"
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "employee.redacted.txt"
    assert redacted["mime"] == "text/plain"

    redacted_text = base64.b64decode(redacted["content_b64"]).decode()
    assert "880412-08-5561" not in redacted_text
    assert "[MYKAD_1]" in redacted_text

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    row = next(x for x in logs["items"] if x["log_id"] == body["log_id"])
    assert row["map_id"] is not None


def test_image_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    image = image_bytes_with_text("Employee IC 880412-08-5561")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "scan.png",
                        "mime": "image/png",
                        "content_b64": base64.b64encode(image).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "scan.redacted.png"
    assert redacted["mime"] == "image/png"

    redacted_doc = extract(redacted["filename"], redacted["mime"], redacted["content_b64"])
    assert "880412-08-5561" not in redacted_doc.text


def test_docx_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    docx = DocxDocument()
    docx.add_paragraph("Employee IC 880412-08-5561 needs a reply.")
    table = docx.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Backup IC 880412-08-5561"
    docx.sections[0].header.paragraphs[0].text = "Header IC 880412-08-5561"
    docx.sections[0].footer.paragraphs[0].text = "Footer IC 880412-08-5561"
    buf = io.BytesIO()
    docx.save(buf)

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "employee.docx",
                        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "content_b64": base64.b64encode(buf.getvalue()).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "employee.redacted.docx"
    assert redacted["mime"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    redacted_doc = DocxDocument(io.BytesIO(base64.b64decode(redacted["content_b64"])))
    redacted_text = "\n".join(
        [p.text for p in redacted_doc.paragraphs]
        + [cell.text for table in redacted_doc.tables for row in table.rows for cell in row.cells]
        + [p.text for section in redacted_doc.sections for p in section.header.paragraphs]
        + [p.text for section in redacted_doc.sections for p in section.footer.paragraphs]
    )
    assert "880412-08-5561" not in redacted_text
    assert "[MYKAD_1]" in redacted_text
    assert "[MYKAD_2]" in redacted_text
    assert "[MYKAD_3]" in redacted_text
    assert "[MYKAD_4]" in redacted_text


def test_docx_attachment_redacts_comments_and_tracked_changes(client: TestClient) -> None:
    docx = docx_bytes_with_comments_and_revisions(
        "Comment IC 880412-08-5561",
        "Inserted IC 880412-08-5561",
        "Deleted IC 880412-08-5561",
    )

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "review.docx",
                        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "content_b64": base64.b64encode(docx).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    redacted_doc = extract(redacted["filename"], redacted["mime"], redacted["content_b64"])

    assert "880412-08-5561" not in redacted_doc.text
    assert "[MYKAD_1]" in redacted_doc.text
    assert "[MYKAD_2]" in redacted_doc.text
    assert "[MYKAD_3]" in redacted_doc.text


def test_docx_attachment_redacts_embedded_ooxml_package(client: TestClient) -> None:
    docx = docx_bytes_with_embedded_docx("Embedded IC 880412-08-5561")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "outer.docx",
                        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "content_b64": base64.b64encode(docx).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    redacted_doc = extract(redacted["filename"], redacted["mime"], redacted["content_b64"])

    assert "880412-08-5561" not in redacted_doc.text
    assert "[MYKAD_1]" in redacted_doc.text


def test_pptx_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Employee IC 880412-08-5561 needs a reply."
    slide.notes_slide.notes_text_frame.text = "Speaker IC 880412-08-5561"
    buf = io.BytesIO()
    deck.save(buf)

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "employee.pptx",
                        "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        "content_b64": base64.b64encode(buf.getvalue()).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "employee.redacted.pptx"
    assert redacted["mime"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    redacted_deck = Presentation(io.BytesIO(base64.b64decode(redacted["content_b64"])))
    redacted_text = "\n".join(
        shape.text
        for slide in redacted_deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text
    )
    redacted_notes = redacted_deck.slides[0].notes_slide.notes_text_frame.text
    redacted_text = "\n".join([redacted_text, redacted_notes])
    assert "880412-08-5561" not in redacted_text
    assert "[MYKAD_1]" in redacted_text
    assert "[MYKAD_2]" in redacted_text


def test_pptx_attachment_redacts_comments(client: TestClient) -> None:
    pptx = pptx_bytes_with_comment("Comment IC 880412-08-5561")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "comments.pptx",
                        "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        "content_b64": base64.b64encode(pptx).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    redacted_doc = extract(redacted["filename"], redacted["mime"], redacted["content_b64"])

    assert "880412-08-5561" not in redacted_doc.text
    assert "[MYKAD_1]" in redacted_doc.text


def test_pptx_attachment_redacts_embedded_ooxml_package(client: TestClient) -> None:
    pptx = pptx_bytes_with_embedded_docx("Embedded IC 880412-08-5561")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "outer.pptx",
                        "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        "content_b64": base64.b64encode(pptx).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    redacted_doc = extract(redacted["filename"], redacted["mime"], redacted["content_b64"])

    assert "880412-08-5561" not in redacted_doc.text
    assert "[MYKAD_1]" in redacted_doc.text


def test_pdf_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    pdf = _pdf_bytes_from_text("Employee IC 880412-08-5561 needs a reply.")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "employee.pdf",
                        "mime": "application/pdf",
                        "content_b64": base64.b64encode(pdf).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "employee.redacted.pdf"
    assert redacted["mime"] == "application/pdf"

    with pdfplumber.open(io.BytesIO(base64.b64decode(redacted["content_b64"]))) as pdf_doc:
        redacted_text = "\n".join(page.extract_text() or "" for page in pdf_doc.pages)
    assert "880412-08-5561" not in redacted_text
    assert "[MYKAD_1]" in redacted_text


def test_scanned_pdf_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    pdf = scanned_pdf_bytes_with_text("Employee IC 880412-08-5561")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "scan.pdf",
                        "mime": "application/pdf",
                        "content_b64": base64.b64encode(pdf).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "scan.redacted.pdf"
    assert redacted["mime"] == "application/pdf"

    with pdfplumber.open(io.BytesIO(base64.b64decode(redacted["content_b64"]))) as pdf_doc:
        redacted_text = "\n".join(page.extract_text() or "" for page in pdf_doc.pages)
    assert "880412-08-5561" not in redacted_text
    assert "[MYKAD_1]" in redacted_text


def test_xlsx_attachment_redaction_returns_redacted_copy(client: TestClient) -> None:
    xlsx = xlsx_bytes_with_shared_strings("Employee IC 880412-08-5561 needs a reply.")

    r = inspect(client, user_id="aiman.hr@acme.my",
                prompt="see attached", attachments=[
                    {
                        "filename": "employee.xlsx",
                        "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        "content_b64": base64.b64encode(xlsx).decode(),
                    },
                ])

    body = r.json()
    assert body["decision"] == "redact"
    assert "MYKAD" in body["data_classes"]
    assert len(body["redacted_attachments"]) == 1

    redacted = body["redacted_attachments"][0]
    assert redacted["filename"] == "employee.redacted.xlsx"
    assert redacted["mime"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    redacted_doc = extract(
        redacted["filename"],
        redacted["mime"],
        redacted["content_b64"],
    )
    assert "880412-08-5561" not in redacted_doc.text
    assert "[MYKAD_1]" in redacted_doc.text


# --------------------------------------------------------------------------
# Two-pass "state & allow" purpose flow, and the hold-timeout downgrade.
# --------------------------------------------------------------------------


def test_two_pass_purpose_flow_returns_require_purpose_then_resolves_on_resubmit(
    client: TestClient, monkeypatch: pytest.MonkeyPatch,
) -> None:
    # M2 classes (e.g. SALARY) no longer carry a require_purpose gate --
    # DEPT_CONFIDENTIAL's allow-band tier is now the only rule left that
    # demonstrates the "state & allow" two-pass flow. 0.78 clears both
    # M8_SPAN_THRESHOLD (0.72, needed for a span/data_class to exist at
    # all) and lands inside the embed allow-band (0.50-0.84).
    seed_internal_doc_for_m8(client, monkeypatch, department="hr", similarity=0.78)
    prompt = ("Please summarize the HR confidential termination plan "
              "for employee grievance cases.")

    pass1 = inspect(client, user_id="aiman.hr@acme.my", prompt=prompt).json()
    assert pass1["decision"] == "allow"
    assert pass1["require_purpose"] is True

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    pass1_row = next(x for x in logs["items"] if x["log_id"] == pass1["log_id"])
    assert pass1_row["purpose"] == "GENERAL"

    # Pass 2: identical prompt, an attested purpose -- resolves without
    # asking again, and the attested purpose lands on the audit record.
    pass2 = inspect(client, user_id="aiman.hr@acme.my", prompt=prompt,
                    purpose="drafting").json()
    assert pass2["decision"] == "allow"
    assert pass2["require_purpose"] is False

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    pass2_row = next(x for x in logs["items"] if x["log_id"] == pass2["log_id"])
    assert pass2_row["purpose"] == "drafting"


def test_hold_timeout_downgrades_via_monkeypatched_clock(
    client: TestClient, monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed_internal_doc_for_m8_hold(client, monkeypatch)
    monkeypatch.setattr("app.main.HOLD_TIMEOUT_SECONDS", 600)

    base_ts = datetime(2026, 7, 20, 9, 0, tzinfo=timezone.utc)

    class FrozenClock(datetime):
        @classmethod
        def now(cls, tz=None):
            return base_ts

    monkeypatch.setattr("app.main.datetime", FrozenClock)

    r = inspect(
        client, user_id="wei.eng@acme.my",
        prompt="Please summarize the HR confidential termination plan "
               "for employee grievance cases.",
    )
    body = r.json()
    assert body["decision"] == "hold"
    log_id = body["log_id"]

    # Still within the window: no downgrade, no extra key on the wire.
    status = client.get(f"/v1/inspect/{log_id}", headers=HEADERS).json()
    assert status == {"decision": "hold", "status": "pending"}

    class LaterClock(datetime):
        @classmethod
        def now(cls, tz=None):
            return base_ts + timedelta(seconds=601)

    monkeypatch.setattr("app.main.datetime", LaterClock)

    downgraded = client.get(f"/v1/inspect/{log_id}", headers=HEADERS).json()
    assert downgraded["decision"] == "hold"    # the verdict is a historical fact
    assert downgraded["status"] == "approved"
    assert downgraded["redacted_prompt"] is not None

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    row = next(x for x in logs["items"] if x["log_id"] == log_id)
    assert row["status"] == "approved"
    assert row["remark"] == "auto: hold timeout -> redact"


# --------------------------------------------------------------------------
# SEMANTIC_DETECTOR=both -- M8 and M9 in parallel behind CompositeSemanticDetector.
# --------------------------------------------------------------------------


def test_semantic_both_mode_produces_spans_from_both_detectors_and_health_reports_it(
    client: TestClient, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.detect.base import Span

    prompt = "Public webinar details are available. The board forecast is confidential."
    embed_start = prompt.index("Public webinar")
    embed_end = embed_start + len("Public webinar details are available.")
    deberta_start = prompt.index("The board forecast")
    deberta_end = deberta_start + len("The board forecast is confidential.")

    class StubSemanticDetector:
        """Stands in for CompositeSemanticDetector's own output -- this test
        is about main.py's wiring and the audit trail, not re-testing the
        composite's dedup logic (see test_composite.py for that)."""

        def detect(self, doc):
            del doc
            return [
                Span(start=embed_start, end=embed_end, data_class="DEPT_CONFIDENTIAL",
                     detector="embed", confidence=0.90),
                Span(start=deberta_start, end=deberta_end, data_class="DEPT_CONFIDENTIAL",
                     detector="deberta", confidence=0.90),
            ]

        def classify(self, doc):
            del doc
            return "engineering", 0.90

    monkeypatch.setattr("app.main.SEMANTIC_DETECTOR", "both")
    monkeypatch.setattr("app.main.get_embed_detector", lambda internal_docs=None: StubSemanticDetector())

    r = inspect(client, user_id="wei.eng@acme.my", prompt=prompt)
    body = r.json()
    assert "DEPT_CONFIDENTIAL" in body["data_classes"]
    detectors = {d["detector"] for d in body["debug"]["detections"]}
    assert detectors == {"embed", "deberta"}

    logs = client.get("/v1/logs", headers=HEADERS, params={"limit": 200}).json()
    row = next(x for x in logs["items"] if x["log_id"] == body["log_id"])
    assert "DEPT_CONFIDENTIAL" in row["data_classes"]

    health = client.get("/v1/health", headers=HEADERS).json()
    assert health["semantic_detector"] == "both"
    assert health["semantic_submodules"] == ["embed", "deberta"]


# --------------------------------------------------------------------------
# ENV=demo refuses to start with an active mock.
# --------------------------------------------------------------------------


def test_demo_guard_refuses_to_start_with_active_mocks() -> None:
    with pytest.raises(RuntimeError):
        _check_demo_guard("demo", {"M2"})


def test_demo_guard_allows_demo_with_no_mocks_left() -> None:
    _check_demo_guard("demo", set())


def test_demo_guard_allows_dev_with_mocks() -> None:
    _check_demo_guard("dev", {"M2", "M3", "M5", "M7", "M8"})
