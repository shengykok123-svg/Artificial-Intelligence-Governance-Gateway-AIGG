# M1  FastAPI app and routers. Consumes IF-9 (audit, vault, rules, tools, users). Produces IF-2, IF-3, IF-7, IF-10.

from __future__ import annotations

import base64
import binascii
import io
import os
import sys
import textwrap
import uuid
from collections import Counter
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from dotenv import load_dotenv
from docx import Document as DocxDocument
from fastapi import Body, Depends, FastAPI, Header, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import JSONResponse
from pptx import Presentation

from app.detect.base import Classification, Detector, DocumentClassifier, Document, Span
from app.errors import Forbidden
from app.extract.embedded import EmbeddedPackage, iter_embedded_ooxml_packages, replace_embedded_packages
from app.extract.files import UnsupportedFile
from app.extract.docx_ooxml import DocxXmlTextTarget, iter_docx_ooxml_text_targets, redact_docx_ooxml_bytes
from app.extract.ocr import OcrResult, ocr_image_bytes, redact_image_bytes
from app.extract.pptx_ooxml import PptxXmlTextTarget, iter_pptx_ooxml_text_targets, redact_pptx_ooxml_bytes
from app.extract.xlsx import XlsxTextTarget, iter_xlsx_text_targets, redact_xlsx_bytes
from app.policy.engine import PolicyContext, Verdict
from app.repo import Repos, get_repos
from app.repo.base import AuditRecord
from app.schemas import (
    Attachment,
    DetectionDebug,
    InternalDocChunkResponse,
    InternalDocsPage,
    InternalDocUploadRequest,
    InternalDocUploadResponse,
    InspectDebug,
    InspectRequest,
    InspectResponse,
    RedactedAttachment,
)

load_dotenv()

ENV = os.getenv("ENV", "dev")
API_KEY = os.getenv("API_KEY", "dev-key-change-me")
MOCK_MODULES: set[str] = {
    m.strip() for m in os.getenv("MOCK_MODULES", "").split(",")
    if m.strip()
}
SEMANTIC_DETECTOR = os.getenv("SEMANTIC_DETECTOR", "m8").strip().lower()
VALID_DEPARTMENTS = {"finance", "hr", "legal", "engineering"}
INTERNAL_DOC_CHUNK_CHARS = 900
INTERNAL_DOC_CHUNK_OVERLAP = 120
# Anti-shadow-AI: nobody waits unbounded on a hold. 0 disables the downgrade.
HOLD_TIMEOUT_SECONDS = int(os.getenv("HOLD_TIMEOUT_SECONDS", "600"))

# The default admin identity behind the dashboard's hardcoded login. Neither
# POST /v1/logs/{id}/decision nor POST /v1/tools/{host}/decision carries a
# reviewer/actor field on the wire (see schemas.py) -- there is exactly one
# login, and this is it. A real IdP-backed identity is future work.
DASHBOARD_ADMIN = "admin@acme.my"


def _check_demo_guard(env: str, mock_modules: set[str]) -> None:
    """A demo video recorded against mock detection would show redaction
    that never happened. That is not a bug, it is a false claim about the
    product -- refuse to start rather than let it happen by accident."""
    if env == "demo" and mock_modules:
        raise RuntimeError(
            f"ENV=demo refuses to start with MOCK_MODULES={sorted(mock_modules)} "
            "still active. Remove every landed module from MOCK_MODULES "
            "before recording a demo."
        )


_check_demo_guard(ENV, MOCK_MODULES)

print(
    f"[startup] ENV={ENV} semantic_detector={SEMANTIC_DETECTOR} "
    f"MOCK_MODULES={sorted(MOCK_MODULES) or ['none']}",
    file=sys.stderr,
)


# --------------------------------------------------------------------------
# Loaders. Each reads MOCK_MODULES exactly once per call. This is the seam:
# when a module lands, its owner removes the ID here and the mock import
# disappears. Explicit selection, never a try/except ImportError fallback --
# a broken real module must fail loudly, not silently serve a mock.
# --------------------------------------------------------------------------


def get_rules_detector() -> Detector:
    if "M2" in MOCK_MODULES:
        from app.mocks import MockRulesDetector
        return MockRulesDetector()
    from app.detect.rules import RulesDetector
    return RulesDetector()


def get_embed_detector(internal_docs: Any | None = None) -> Detector | DocumentClassifier:
    if SEMANTIC_DETECTOR == "m9_deberta":
        # M9 is opt-in. It has the same two methods as M8, so nothing in the
        # inspect pipeline needs to know which semantic implementation runs.
        from app.detect.deberta import DebertaDetector
        return DebertaDetector()
    if SEMANTIC_DETECTOR == "both":
        # M8 and M9 run in parallel behind the same Detector +
        # DocumentClassifier protocol -- every call site downstream (the
        # inspect pipeline and the attachment-redaction helpers) takes ONE
        # object satisfying it, so this is the entire integration. Fail
        # loudly: DebertaDetector's own missing-weights error is not caught
        # here, and neither is EmbedDetector's missing-repo error.
        from app.detect.composite import CompositeSemanticDetector
        from app.detect.deberta import DebertaDetector
        if "M8" in MOCK_MODULES:
            from app.mocks import MockEmbedDetector
            embed: Detector | DocumentClassifier = MockEmbedDetector()
        else:
            from app.detect.embed import EmbedDetector
            embed = EmbedDetector(internal_docs=internal_docs)
        return CompositeSemanticDetector(embed, DebertaDetector())
    if SEMANTIC_DETECTOR != "m8":
        raise RuntimeError("SEMANTIC_DETECTOR must be 'm8', 'm9_deberta', or 'both'.")
    if "M8" in MOCK_MODULES:
        from app.mocks import MockEmbedDetector
        return MockEmbedDetector()
    from app.detect.embed import EmbedDetector
    return EmbedDetector(internal_docs=internal_docs)


def get_extractor() -> Callable[[str, str, str], Document]:
    from app.extract.files import extract
    return extract


def get_decider() -> Callable[[Classification, PolicyContext, list], Verdict]:
    from app.policy.engine import decide
    return decide


def get_tokeniser(vault: Any) -> Any:
    if "M7" in MOCK_MODULES:
        from app.mocks import MockTokeniser
        return MockTokeniser(vault)
    from app.vault.tokeniser import Tokeniser
    return Tokeniser(vault)


def get_m8_encoder() -> Any:
    from app.detect.embed import MODEL_NAME
    from sentence_transformers import SentenceTransformer

    return SentenceTransformer(MODEL_NAME, local_files_only=True)


def _chunk_internal_doc_text(text: str) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines())
    normalized = "\n".join(line for line in normalized.splitlines() if line)
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + INTERNAL_DOC_CHUNK_CHARS, len(normalized))
        if end < len(normalized):
            break_at = max(
                normalized.rfind("\n", start, end),
                normalized.rfind(". ", start, end),
                normalized.rfind(" ", start, end),
            )
            if break_at > start + 200:
                end = break_at + 1
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - INTERNAL_DOC_CHUNK_OVERLAP, start + 1)
    return chunks


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PDF_MIME = "application/pdf"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IMAGE_MIMES = {"image/png", "image/jpeg"}


def _can_redact_text_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    filename = att.filename.lower()
    return (
        normalised_mime in {"text/plain", "text/csv"}
        or filename.endswith(".txt")
        or filename.endswith(".csv")
    )


def _can_redact_docx_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    return normalised_mime == DOCX_MIME or att.filename.lower().endswith(".docx")


def _can_redact_pptx_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    return normalised_mime == PPTX_MIME or att.filename.lower().endswith(".pptx")


def _can_redact_pdf_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    return normalised_mime == PDF_MIME or att.filename.lower().endswith(".pdf")


def _can_redact_image_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    return normalised_mime in IMAGE_MIMES or att.filename.lower().endswith((".png", ".jpg", ".jpeg"))


def _can_redact_xlsx_attachment(att: Attachment) -> bool:
    normalised_mime = att.mime.lower().split(";")[0].strip()
    return normalised_mime == XLSX_MIME or att.filename.lower().endswith(".xlsx")


def _redacted_filename(filename: str) -> str:
    if "." not in filename:
        return f"{filename}.redacted.txt"
    stem, ext = filename.rsplit(".", 1)
    return f"{stem}.redacted.{ext}"


def _redacted_image_filename(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    return f"{stem}.redacted.png"


def _redacted_attachment(att: Attachment, redacted_text: str) -> RedactedAttachment:
    return RedactedAttachment(
        filename=_redacted_filename(att.filename),
        mime=att.mime,
        content_b64=base64.b64encode(redacted_text.encode("utf-8")).decode("ascii"),
    )


def _decode_attachment_bytes(att: Attachment) -> bytes:
    try:
        return base64.b64decode(att.content_b64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise UnsupportedFile(f"invalid base64 (filename={att.filename!r})") from exc


def _iter_docx_text_targets(docx: Any) -> list[Any]:
    targets: list[Any] = []
    for container in _iter_docx_story_containers(docx):
        targets.extend(p for p in container.paragraphs if p.text)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        targets.append(cell)
    return targets


def _iter_docx_story_containers(docx: Any) -> list[Any]:
    containers: list[Any] = [docx]
    seen = {id(docx._element)}
    for section in docx.sections:
        for attr in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            container = getattr(section, attr)
            marker = id(container._element)
            if marker in seen:
                continue
            seen.add(marker)
            containers.append(container)
    return containers


def _redacted_docx_attachment(
    att: Attachment,
    docx: Any,
    xml_replacements: dict[DocxXmlTextTarget, str] | None = None,
    embedded_replacements: dict[str, bytes] | None = None,
) -> RedactedAttachment:
    buf = io.BytesIO()
    docx.save(buf)
    content = buf.getvalue()
    if xml_replacements:
        content = redact_docx_ooxml_bytes(content, xml_replacements)
    if embedded_replacements:
        content = replace_embedded_packages(content, embedded_replacements)
    return RedactedAttachment(
        filename=_redacted_filename(att.filename),
        mime=att.mime or DOCX_MIME,
        content_b64=base64.b64encode(content).decode("ascii"),
    )


def _iter_pptx_text_targets(deck: Any) -> list[Any]:
    targets: list[Any] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text_frame = shape.text_frame
            if text_frame.text:
                targets.append(text_frame)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            targets.append(slide.notes_slide.notes_text_frame)
    return targets


def _redacted_pptx_attachment(
    att: Attachment,
    deck: Any,
    xml_replacements: dict[PptxXmlTextTarget, str] | None = None,
    embedded_replacements: dict[str, bytes] | None = None,
) -> RedactedAttachment:
    buf = io.BytesIO()
    deck.save(buf)
    content = buf.getvalue()
    if xml_replacements:
        content = redact_pptx_ooxml_bytes(content, xml_replacements)
    if embedded_replacements:
        content = replace_embedded_packages(content, embedded_replacements)
    return RedactedAttachment(
        filename=_redacted_filename(att.filename),
        mime=att.mime or PPTX_MIME,
        content_b64=base64.b64encode(content).decode("ascii"),
    )


def _pdf_literal_string(text: str) -> bytes:
    raw = text.encode("latin-1", errors="replace")
    escaped = bytearray()
    for byte in raw:
        if byte in {0x28, 0x29, 0x5C}:  # (, ), backslash
            escaped.append(0x5C)
            escaped.append(byte)
        elif byte == 0x0D:
            escaped.extend(b"\\r")
        elif byte == 0x0A:
            escaped.extend(b"\\n")
        elif byte == 0x09:
            escaped.extend(b"\\t")
        else:
            escaped.append(byte)
    return b"(" + bytes(escaped) + b")"


def _pdf_bytes_from_text(text: str) -> bytes:
    wrapped: list[str] = []
    for raw_line in text.splitlines() or [""]:
        wrapped.extend(textwrap.wrap(raw_line, width=86) or [""])

    lines_per_page = 42
    pages = [wrapped[i:i + lines_per_page] for i in range(0, len(wrapped), lines_per_page)] or [[""]]

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    page_ids: list[int] = []

    for page_lines in pages:
        page_id = len(objects) + 1
        content_id = page_id + 1
        page_ids.append(page_id)

        content_lines = [b"BT", b"/F1 11 Tf", b"72 740 Td", b"14 TL"]
        for idx, line in enumerate(page_lines):
            if idx:
                content_lines.append(b"T*")
            content_lines.append(_pdf_literal_string(line) + b" Tj")
        content_lines.append(b"ET")
        content = b"\n".join(content_lines)

        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode("ascii")
        )
        objects.append(
            b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream"
        )

    kids = b" ".join(f"{page_id} 0 R".encode("ascii") for page_id in page_ids)
    objects[1] = b"<< /Type /Pages /Count " + str(len(page_ids)).encode("ascii") + b" /Kids [" + kids + b"] >>"

    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for obj_id, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{obj_id} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def _redacted_pdf_attachment(att: Attachment, redacted_text: str) -> RedactedAttachment:
    return RedactedAttachment(
        filename=_redacted_filename(att.filename),
        mime=att.mime or PDF_MIME,
        content_b64=base64.b64encode(_pdf_bytes_from_text(redacted_text)).decode("ascii"),
    )


def _redacted_image_attachment(att: Attachment, result: OcrResult, spans: list[Span]) -> RedactedAttachment:
    redacted_bytes = redact_image_bytes(
        _decode_attachment_bytes(att),
        result,
        [(span.start, span.end) for span in spans],
    )
    return RedactedAttachment(
        filename=_redacted_image_filename(att.filename),
        mime="image/png",
        content_b64=base64.b64encode(redacted_bytes).decode("ascii"),
    )


def _redacted_xlsx_attachment(att: Attachment, replacements: dict[XlsxTextTarget, str]) -> RedactedAttachment:
    redacted_bytes = redact_xlsx_bytes(_decode_attachment_bytes(att), replacements)
    return RedactedAttachment(
        filename=_redacted_filename(att.filename),
        mime=att.mime or XLSX_MIME,
        content_b64=base64.b64encode(redacted_bytes).decode("ascii"),
    )


def _embedded_replacements_for(
    raw: bytes,
    prefix: str,
    rules_detector: Detector,
    embed_detector: Detector | DocumentClassifier,
    tokeniser: Any,
) -> tuple[dict[str, bytes], str | None]:
    replacements: dict[str, bytes] = {}
    map_id: str | None = None
    for package in iter_embedded_ooxml_packages(raw, prefix):
        redacted = _redacted_embedded_package(package, rules_detector, embed_detector, tokeniser)
        if redacted is None:
            continue
        replacements[package.part_name] = base64.b64decode(redacted.content_b64)
        map_id = map_id or "embedded"
    return replacements, map_id


def _redacted_embedded_package(
    package: EmbeddedPackage,
    rules_detector: Detector,
    embed_detector: Detector | DocumentClassifier,
    tokeniser: Any,
) -> RedactedAttachment | None:
    att = Attachment(
        filename=package.filename,
        mime=package.mime,
        content_b64=base64.b64encode(package.raw).decode("ascii"),
    )
    doc = get_extractor()(att.filename, att.mime, att.content_b64)
    spans = rules_detector.detect(doc) + embed_detector.detect(doc)
    if not spans:
        return None

    tokenise_inputs: list[tuple[str, list[Span]]] = [("", []), (doc.text, spans)]
    redacted_items, _ = tokeniser.tokenise_many(tokenise_inputs)

    if _can_redact_docx_attachment(att):
        docx = DocxDocument(io.BytesIO(package.raw))
        xml_replacements = _docx_xml_replacements(package.raw, att, rules_detector, embed_detector, tokeniser)
        for target in _iter_docx_text_targets(docx):
            target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
            target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
            if target_spans:
                redacted_target, _ = tokeniser.tokenise(target.text, target_spans)
                target.text = redacted_target
        return _redacted_docx_attachment(att, docx, xml_replacements)

    if _can_redact_pptx_attachment(att):
        deck = Presentation(io.BytesIO(package.raw))
        xml_replacements = _pptx_xml_replacements(package.raw, att, rules_detector, embed_detector, tokeniser)
        for target in _iter_pptx_text_targets(deck):
            target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
            target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
            if target_spans:
                redacted_target, _ = tokeniser.tokenise(target.text, target_spans)
                target.text = redacted_target
        return _redacted_pptx_attachment(att, deck, xml_replacements)

    if _can_redact_xlsx_attachment(att):
        replacements: dict[XlsxTextTarget, str] = {}
        for target in iter_xlsx_text_targets(package.raw):
            target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
            target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
            if target_spans:
                redacted_target, _ = tokeniser.tokenise(target.text, target_spans)
                replacements[target] = redacted_target
        return _redacted_xlsx_attachment(att, replacements) if replacements else None

    return _redacted_attachment(att, redacted_items[1])


def _docx_xml_replacements(
    raw: bytes,
    att: Attachment,
    rules_detector: Detector,
    embed_detector: Detector | DocumentClassifier,
    tokeniser: Any,
) -> dict[DocxXmlTextTarget, str]:
    replacements: dict[DocxXmlTextTarget, str] = {}
    for target in iter_docx_ooxml_text_targets(raw):
        target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
        target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
        if target_spans:
            redacted_target, _ = tokeniser.tokenise(target.text, target_spans)
            replacements[target] = redacted_target
    return replacements


def _pptx_xml_replacements(
    raw: bytes,
    att: Attachment,
    rules_detector: Detector,
    embed_detector: Detector | DocumentClassifier,
    tokeniser: Any,
) -> dict[PptxXmlTextTarget, str]:
    replacements: dict[PptxXmlTextTarget, str] = {}
    for target in iter_pptx_ooxml_text_targets(raw):
        target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
        target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
        if target_spans:
            redacted_target, _ = tokeniser.tokenise(target.text, target_spans)
            replacements[target] = redacted_target
    return replacements


# --------------------------------------------------------------------------
# App
# --------------------------------------------------------------------------


def require_api_key(x_api_key: str | None = Header(None)) -> None:
    if x_api_key != API_KEY:
        raise HTTPException(status_code=403, detail="invalid API key")


app = FastAPI(title="AI Governance Gateway",
              dependencies=[Depends(require_api_key)])

# Local-only testing UI. It makes its own authenticated API call from the
# browser, so the page contains no bypass around the API-key dependency.
app.mount("/m9-tester", StaticFiles(directory=Path(__file__).parent / "static", html=True), name="m9-tester")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_origin_regex=r"chrome-extension://.*",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(UnsupportedFile)
def _unsupported_file_handler(request: Request, exc: UnsupportedFile) -> JSONResponse:
    return JSONResponse(status_code=415, content={"detail": str(exc)})


@app.exception_handler(Forbidden)
def _forbidden_handler(request: Request, exc: Forbidden) -> JSONResponse:
    return JSONResponse(status_code=403, content={"detail": str(exc) or "forbidden"})


def check_m8_index_not_empty() -> None:
    """M8's index has no reload signal: an empty internal_doc_chunks table
    means the real (unmocked) EmbedDetector will simply never produce a
    span or a department -- not an error, just silent. Legitimate for a
    fresh memory-backend dev environment; worth a loud warning anywhere
    else, since nothing else would ever surface it. Warning, not error --
    unlike a missing M9 weights file, an empty index is not a broken
    module, just an unseeded one."""
    if SEMANTIC_DETECTOR not in ("m8", "both") or "M8" in MOCK_MODULES:
        return
    from app.db import REPO_BACKEND, SessionLocal
    if REPO_BACKEND == "postgres":
        session = SessionLocal()
        try:
            from app.repo.postgres import PgInternalDocRepo
            _rows, total = PgInternalDocRepo(session).list(limit=1, offset=0)
        finally:
            session.close()
    else:
        from app.repo import memory_repos
        _rows, total = memory_repos().internal_docs.list(limit=1, offset=0)
    if total == 0:
        print(
            "[startup] WARNING: internal_doc_chunks is empty -- M8 will "
            "never produce a span or department until documents are "
            "indexed via POST /v1/internal-docs.",
            file=sys.stderr,
        )


@app.on_event("startup")
def _startup_checks() -> None:
    check_m8_index_not_empty()


# --------------------------------------------------------------------------
# Pipeline
# --------------------------------------------------------------------------


@app.post("/v1/inspect", response_model=InspectResponse)
def inspect(req: InspectRequest, repos: Repos = Depends(get_repos)) -> InspectResponse:
    # SD-2: role and department come from the users table, keyed by user_id.
    # Never a header -- a self-asserted role is a self-issued permission.
    user = repos.users.by_id(req.user_id)
    if user is None:
        raise Forbidden(f"unknown user_id {req.user_id!r}")

    # SD-7: an unknown host auto-registers as PUBLIC_AI/unapproved and is
    # never held on that account alone -- this is the shadow-AI story.
    tool = repos.tools.by_host(req.target_host)
    if tool is None:
        tool = repos.tools.register(req.target_host)

    extract = get_extractor()
    documents: list[Document] = [
        Document(text=req.prompt, source_filename=None, mime="text/plain")
    ]
    attachment_docs: list[tuple[Attachment, Document]] = []
    for att in req.attachments:
        doc = extract(att.filename, att.mime, att.content_b64)
        documents.append(doc)
        attachment_docs.append((att, doc))

    rules_detector = get_rules_detector()
    embed_detector = get_embed_detector(repos.internal_docs)

    # Span.start/end are offsets into ONE Document's text. The prompt's own
    # spans are what get redacted back into the response (redacted_prompt is
    # a single field, and so is AuditRecord.prompt_text); attachment spans
    # still count toward the verdict via data_classes, but are never used to
    # slice the prompt string.
    prompt_doc = documents[0]
    prompt_spans = rules_detector.detect(prompt_doc) + embed_detector.detect(prompt_doc)

    all_spans: list[Span] = list(prompt_spans)
    debug_detections = [
        DetectionDebug(
            data_class=span.data_class,
            detector=span.detector,
            confidence=span.confidence,
            source="prompt",
        )
        for span in prompt_spans
    ]
    attachment_spans: list[tuple[Attachment, Document, list[Span]]] = []
    for att, doc in attachment_docs:
        spans = rules_detector.detect(doc) + embed_detector.detect(doc)
        all_spans += spans
        attachment_spans.append((att, doc, spans))
        debug_detections.extend(
            DetectionDebug(
                data_class=span.data_class,
                detector=span.detector,
                confidence=span.confidence,
                source="attachment",
            )
            for span in spans
        )

    data_classes = sorted({span.data_class for span in all_spans})
    confidence = max((s.confidence for s in all_spans), default=0.0)

    combined_text = "\n".join(doc.text for doc in documents)
    classify_doc = Document(text=combined_text, source_filename=None, mime="text/plain")

    # M10: CompositeSemanticDetector exposes classify_debug() so the audit
    # trail can carry M8's and M9's raw opinions alongside the fused one,
    # without running either model twice. Anything that only implements the
    # plain classify() (a lone EmbedDetector/DebertaDetector, a mock, a test
    # stub) falls back to attributing that single result to whichever
    # detector SEMANTIC_DETECTOR says is actually running.
    classify_debug = getattr(embed_detector, "classify_debug", None)
    if classify_debug is not None:
        (content_department, department_confidence), m8_raw, m9_raw, m9_sensitivity_confidence = classify_debug(classify_doc)
        m8_department, m8_confidence = m8_raw
        m9_department, m9_confidence = m9_raw
    else:
        content_department, department_confidence = embed_detector.classify(classify_doc)
        classify_sensitivity = getattr(embed_detector, "classify_sensitivity", None)
        if SEMANTIC_DETECTOR == "m9_deberta":
            m8_department, m8_confidence = None, 0.0
            m9_department, m9_confidence = content_department, department_confidence
            m9_sensitivity_confidence = classify_sensitivity(classify_doc) if classify_sensitivity is not None else 0.0
        else:
            m8_department, m8_confidence = content_department, department_confidence
            m9_department, m9_confidence = None, 0.0
            m9_sensitivity_confidence = 0.0

    cross_department_confidence: float | None = None
    if content_department is not None and content_department != user.department:
        cross_department_confidence = max(m8_confidence, m9_confidence)

    # M10 fix: decide()'s per-detector confidence gate (embed/deberta) only
    # has NO classification.confidence fallback -- unlike "*" rules, it looks
    # exclusively at classification.spans for that (class, detector) pair. If
    # spans here were prompt-only, a DEPT_CONFIDENTIAL span found ONLY in an
    # attachment (no echo in the prompt text) would be invisible to the tier
    # table and silently default to allow. all_spans (prompt + attachment)
    # fixes that; spans_to_redact is still filtered back down to prompt-only
    # spans below before it ever touches req.prompt, so mixing attachment
    # spans in here cannot corrupt prompt-text slicing.
    classification = Classification(
        data_classes=data_classes,
        confidence=confidence,
        spans=all_spans,
        content_department=content_department,
        department_confidence=department_confidence,
    )

    ctx = PolicyContext(
        user_role=user.role,
        department=user.department,
        ai_tool_tier=tool.tier,
        tool_status=tool.status,
        purpose=req.purpose,
        target_host=req.target_host,
    )

    decide = get_decider()
    verdict = decide(classification, ctx, repos.rules.all())

    redacted_prompt: str | None = None
    redacted_attachments: list[RedactedAttachment] = []
    map_id: str | None = None
    prompt_text_for_audit = req.prompt

    # Hold and deny: prepare the redaction eagerly too, not just for redact.
    # Hold needs it so the lazy timeout downgrade in GET /v1/inspect/{log_id}
    # never has to run M7 itself. Deny needs it because SD-1 promises the
    # audit row never holds the raw original -- a denied request still gets
    # logged, and whatever M2 found (a credential, a MyKad number) must be
    # masked there even though the request itself is blocked outright and
    # nothing is ever released to the caller. Nothing is exposed to the
    # caller yet on either path -- decision stays hold/deny and
    # computed_redacted_prompt/attachments are withheld from the response
    # below regardless.
    if verdict.decision in ("redact", "hold", "deny"):
        computed_redacted_attachments: list[RedactedAttachment] = []
        tokeniser = get_tokeniser(repos.vault)
        # classification.spans (and therefore verdict.spans_to_redact) can
        # now contain attachment-sourced spans too (see the Classification
        # construction above) -- intersect back down to prompt_spans before
        # slicing req.prompt, so an attachment's offsets can never be sliced
        # against a different document's text.
        prompt_span_set = set(prompt_spans)
        prompt_spans_to_redact = [s for s in verdict.spans_to_redact if s in prompt_span_set]
        tokenise_inputs: list[tuple[str, list[Span]]] = [(req.prompt, prompt_spans_to_redact)]
        attachment_redaction_inputs: list[tuple[Attachment, int]] = []
        docx_redaction_inputs: list[
            tuple[Attachment, Any, list[tuple[Any, int]], list[tuple[DocxXmlTextTarget, int]], dict[str, bytes]]
        ] = []
        pptx_redaction_inputs: list[
            tuple[Attachment, Any, list[tuple[Any, int]], list[tuple[PptxXmlTextTarget, int]], dict[str, bytes]]
        ] = []
        pdf_redaction_inputs: list[tuple[Attachment, int]] = []
        image_redaction_inputs: list[tuple[Attachment, OcrResult, list[Span]]] = []
        xlsx_redaction_inputs: list[tuple[Attachment, list[tuple[XlsxTextTarget, int]]]] = []
        for att, doc, spans in attachment_spans:
            if not spans:
                continue
            if _can_redact_text_attachment(att):
                attachment_redaction_inputs.append((att, len(tokenise_inputs)))
                tokenise_inputs.append((doc.text, spans))
                continue

            if _can_redact_docx_attachment(att):
                raw = _decode_attachment_bytes(att)
                docx = DocxDocument(io.BytesIO(raw))
                targets: list[tuple[Any, int]] = []
                for target in _iter_docx_text_targets(docx):
                    target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
                    target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                    if not target_spans:
                        continue
                    targets.append((target, len(tokenise_inputs)))
                    tokenise_inputs.append((target.text, target_spans))
                xml_targets: list[tuple[DocxXmlTextTarget, int]] = []
                for target in iter_docx_ooxml_text_targets(raw):
                    target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
                    target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                    if not target_spans:
                        continue
                    xml_targets.append((target, len(tokenise_inputs)))
                    tokenise_inputs.append((target.text, target_spans))
                embedded_replacements, embedded_map_id = _embedded_replacements_for(
                    raw, "word/embeddings/", rules_detector, embed_detector, tokeniser
                )
                if embedded_map_id:
                    map_id = map_id or embedded_map_id
                if targets or xml_targets or embedded_replacements:
                    docx_redaction_inputs.append((att, docx, targets, xml_targets, embedded_replacements))
                continue

            if _can_redact_pptx_attachment(att):
                raw = _decode_attachment_bytes(att)
                deck = Presentation(io.BytesIO(raw))
                targets: list[tuple[Any, int]] = []
                for target in _iter_pptx_text_targets(deck):
                    target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
                    target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                    if not target_spans:
                        continue
                    targets.append((target, len(tokenise_inputs)))
                    tokenise_inputs.append((target.text, target_spans))
                xml_targets: list[tuple[PptxXmlTextTarget, int]] = []
                for target in iter_pptx_ooxml_text_targets(raw):
                    target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
                    target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                    if not target_spans:
                        continue
                    xml_targets.append((target, len(tokenise_inputs)))
                    tokenise_inputs.append((target.text, target_spans))
                embedded_replacements, embedded_map_id = _embedded_replacements_for(
                    raw, "ppt/embeddings/", rules_detector, embed_detector, tokeniser
                )
                if embedded_map_id:
                    map_id = map_id or embedded_map_id
                if targets or xml_targets or embedded_replacements:
                    pptx_redaction_inputs.append((att, deck, targets, xml_targets, embedded_replacements))
                continue

            if _can_redact_pdf_attachment(att):
                pdf_redaction_inputs.append((att, len(tokenise_inputs)))
                tokenise_inputs.append((doc.text, spans))
                continue

            if _can_redact_image_attachment(att):
                result = ocr_image_bytes(_decode_attachment_bytes(att))
                target_doc = Document(text=result.text, source_filename=att.filename, mime=att.mime)
                target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                if target_spans:
                    image_redaction_inputs.append((att, result, target_spans))
                    tokenise_inputs.append((result.text, target_spans))
                continue

            if _can_redact_xlsx_attachment(att):
                targets: list[tuple[XlsxTextTarget, int]] = []
                for target in iter_xlsx_text_targets(_decode_attachment_bytes(att)):
                    target_doc = Document(text=target.text, source_filename=att.filename, mime=att.mime)
                    target_spans = rules_detector.detect(target_doc) + embed_detector.detect(target_doc)
                    if not target_spans:
                        continue
                    targets.append((target, len(tokenise_inputs)))
                    tokenise_inputs.append((target.text, target_spans))
                if targets:
                    xlsx_redaction_inputs.append((att, targets))

        redacted_items, map_id = tokeniser.tokenise_many(tokenise_inputs)
        computed_redacted_prompt = redacted_items[0]
        for att, redacted_index in attachment_redaction_inputs:
            computed_redacted_attachments.append(_redacted_attachment(att, redacted_items[redacted_index]))
        for att, docx, targets, xml_targets, embedded_replacements in docx_redaction_inputs:
            for target, redacted_index in targets:
                target.text = redacted_items[redacted_index]
            xml_replacements = {target: redacted_items[redacted_index] for target, redacted_index in xml_targets}
            computed_redacted_attachments.append(_redacted_docx_attachment(att, docx, xml_replacements, embedded_replacements))
        for att, deck, targets, xml_targets, embedded_replacements in pptx_redaction_inputs:
            for target, redacted_index in targets:
                target.text = redacted_items[redacted_index]
            xml_replacements = {target: redacted_items[redacted_index] for target, redacted_index in xml_targets}
            computed_redacted_attachments.append(_redacted_pptx_attachment(att, deck, xml_replacements, embedded_replacements))
        for att, redacted_index in pdf_redaction_inputs:
            computed_redacted_attachments.append(_redacted_pdf_attachment(att, redacted_items[redacted_index]))
        for att, result, target_spans in image_redaction_inputs:
            computed_redacted_attachments.append(_redacted_image_attachment(att, result, target_spans))
        for att, targets in xlsx_redaction_inputs:
            replacements = {target: redacted_items[redacted_index] for target, redacted_index in targets}
            computed_redacted_attachments.append(_redacted_xlsx_attachment(att, replacements))

        # SD-1: the audit row holds the REDACTED text; the original lives
        # only in the vault. For hold, this is a PREPARED redaction: the
        # response below still withholds it (decision stays "hold") until
        # either an admin approves or the lazy timeout in
        # GET /v1/inspect/{log_id} releases it.
        prompt_text_for_audit = computed_redacted_prompt
        if verdict.decision == "redact":
            redacted_prompt = computed_redacted_prompt
            redacted_attachments = computed_redacted_attachments

    status = "pending" if verdict.decision == "hold" else "auto"

    log_id = repos.audit.create(AuditRecord(
        log_id="",
        user_id=req.user_id,
        ts=datetime.now(timezone.utc),
        target_host=req.target_host,
        prompt_text=prompt_text_for_audit,
        filename=req.attachments[0].filename if req.attachments else None,
        data_classes=list(data_classes),
        decision=verdict.decision,
        reason=verdict.reason,
        status=status,
        reviewer=None,
        reviewed_at=None,
        remark=None,
        map_id=map_id,
        user_department=user.department,
        content_department=content_department,
        department_confidence=department_confidence,
        purpose=req.purpose,
        m8_department=m8_department,
        m8_confidence=m8_confidence,
        m9_department=m9_department,
        m9_confidence=m9_confidence,
        m9_sensitivity_confidence=m9_sensitivity_confidence,
        cross_department_confidence=cross_department_confidence,
    ))

    return InspectResponse(
        decision=verdict.decision,       # type: ignore[arg-type]
        log_id=log_id,
        data_classes=list(data_classes),
        redacted_prompt=redacted_prompt,
        redacted_attachments=redacted_attachments,
        reason=verdict.reason,
        require_purpose=verdict.require_purpose,
        content_department=content_department,
        department_confidence=department_confidence,
        m8_department=m8_department,
        m8_confidence=m8_confidence,
        m9_department=m9_department,
        m9_confidence=m9_confidence,
        m9_sensitivity_confidence=m9_sensitivity_confidence,
        debug=InspectDebug(
            policy_layer="M3",
            semantic_detector=SEMANTIC_DETECTOR,
            detections=debug_detections,
        ),
    )


# --------------------------------------------------------------------------
# Audit / approval loop
# --------------------------------------------------------------------------


@app.get("/v1/inspect/{log_id}")
def get_inspect_status(log_id: str, repos: Repos = Depends(get_repos)) -> dict:
    rec = repos.audit.get(log_id)
    if rec is None:
        raise HTTPException(status_code=404, detail="unknown log_id")

    # Lazy timeout downgrade: no scheduler, evaluated the next time anyone
    # asks. prompt_text already holds the PREPARED redaction from inspect
    # time (see inspect()) -- nobody waits unbounded on a human reviewer.
    if (
        rec.status == "pending"
        and HOLD_TIMEOUT_SECONDS > 0
        and (datetime.now(timezone.utc) - rec.ts).total_seconds() > HOLD_TIMEOUT_SECONDS
    ):
        rec = repos.audit.set_decision(
            log_id, approve=True, reviewer=DASHBOARD_ADMIN,
            remark="auto: hold timeout -> redact",
        )

    # rec.prompt_text always holds the PREPARED redaction (see inspect()),
    # whether this hold was released by a human reviewer approving it in the
    # dashboard or by the lazy timeout above -- a regex-caught credential or
    # PII span is masked either way; a reviewer approving the DEPT_CONFIDENTIAL
    # judgement was never asked to also bless an unrelated leaked secret.
    # Only "approved" ever has something to forward; "pending"/"denied" don't.
    if rec.status == "approved":
        return {
            "decision": rec.decision,
            "status": rec.status,
            "redacted_prompt": rec.prompt_text,
        }

    return {"decision": rec.decision, "status": rec.status}


@app.get("/v1/logs")
def list_logs(
    limit: int = Query(20, ge=1, le=200),
    offset: int = Query(0, ge=0),
    status: str | None = None,
    decision: str | None = None,
    cross_department_only: bool = Query(False),
    repos: Repos = Depends(get_repos),
) -> dict:
    rows, total = repos.audit.list(limit=limit, offset=offset,
                                   status=status, decision=decision,
                                   cross_department_only=cross_department_only)
    return {"items": rows, "total": total, "limit": limit, "offset": offset}


@app.post("/v1/logs/{log_id}/decision")
def decide_log(
    log_id: str,
    approve: bool = Body(...),
    remark: str = Body(...),
    repos: Repos = Depends(get_repos),
) -> AuditRecord:
    try:
        return repos.audit.set_decision(
            log_id, approve=approve, reviewer=DASHBOARD_ADMIN, remark=remark,
        )
    except KeyError:
        raise HTTPException(status_code=404, detail="unknown log_id")


@app.get("/v1/stats")
def stats(repos: Repos = Depends(get_repos)):
    return repos.audit.stats()


# --------------------------------------------------------------------------
# Policy and tools (IF-10)
# --------------------------------------------------------------------------


@app.get("/v1/rules")
def list_rules(repos: Repos = Depends(get_repos)):
    return repos.rules.all()


@app.get("/v1/tools")
def list_tools(repos: Repos = Depends(get_repos)) -> list[dict]:
    # TODO(team decision): no repo method returns a usage count. Stopgap --
    # page repos.audit.list() and count target_host in Python. At 30 seed
    # rows this is free; it will not be at 30,000.
    all_rows, _total = repos.audit.list(limit=10_000, offset=0)
    usage = Counter(row.target_host for row in all_rows)
    return [
        {**asdict(tool), "usage_count": usage.get(tool.host, 0)}
        for tool in repos.tools.all()
    ]


@app.post("/v1/tools/{host}/decision")
def decide_tool(
    host: str,
    status: str = Body(...),
    remark: str = Body(...),
    repos: Repos = Depends(get_repos),
):
    if repos.tools.by_host(host) is None:
        raise HTTPException(status_code=404, detail="unknown host")
    # _set_status is not on the ToolRepo protocol yet -- see the note in
    # app/repo/memory.py. Both backends implement it; M1 wires it here until
    # the team promotes it to IF-9. AiTool has no remark column (open
    # question, see docs/prompts/M1.md notes) -- accepted on the wire,
    # nowhere to persist it yet.
    del remark
    return repos.tools._set_status(host, status, DASHBOARD_ADMIN)  # type: ignore[attr-defined]


# --------------------------------------------------------------------------
# M8 internal-document corpus
# --------------------------------------------------------------------------


@app.get("/v1/internal-docs", response_model=InternalDocsPage)
def list_internal_docs(
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
    repos: Repos = Depends(get_repos),
) -> InternalDocsPage:
    rows, total = repos.internal_docs.list(limit=limit, offset=offset)
    return InternalDocsPage(
        items=[
            InternalDocChunkResponse(
                id=row.id,
                source_file=row.source_file,
                department=row.department,
                chunk_text=row.chunk_text,
                created_at=row.created_at.isoformat(),
            )
            for row in rows
        ],
        total=total,
        limit=limit,
        offset=offset,
    )


@app.post("/v1/internal-docs", response_model=InternalDocUploadResponse)
def upload_internal_doc(
    req: InternalDocUploadRequest,
    repos: Repos = Depends(get_repos),
) -> InternalDocUploadResponse:
    department = req.department.strip().lower()
    if department not in VALID_DEPARTMENTS:
        raise HTTPException(status_code=422, detail="department must be finance, hr, legal, or engineering")

    doc = get_extractor()(req.attachment.filename, req.attachment.mime, req.attachment.content_b64)
    chunks = _chunk_internal_doc_text(doc.text)
    if not chunks:
        raise HTTPException(status_code=422, detail="no text could be extracted from this internal document")

    encoder = get_m8_encoder()
    embeddings = encoder.encode(chunks, normalize_embeddings=True)

    if req.replace_source:
        repos.internal_docs.clear_source(req.attachment.filename)
    chunk_count = repos.internal_docs.add_chunks(
        [
            (str(uuid.uuid4()), chunk, [float(value) for value in embedding.tolist()])
            for chunk, embedding in zip(chunks, embeddings)
        ],
        source_file=req.attachment.filename,
        department=department,
    )
    _rows, total = repos.internal_docs.list(limit=1, offset=0)
    return InternalDocUploadResponse(
        source_file=req.attachment.filename,
        department=department,
        chunk_count=chunk_count,
        total=total,
    )


# --------------------------------------------------------------------------
# Health
# --------------------------------------------------------------------------


@app.get("/v1/health")
def health() -> dict:
    body: dict[str, Any] = {
        "status": "ok",
        "semantic_detector": SEMANTIC_DETECTOR,
        "mocked": sorted(MOCK_MODULES),
    }
    if SEMANTIC_DETECTOR == "both":
        # "both" is not itself a detector -- name the two that composed it,
        # so a caller polling health can tell M8+M9 are the effective mode
        # without knowing to interpret the string "both" specially.
        body["semantic_submodules"] = ["embed", "deberta"]
    return body
