"""M5 file text extraction.

extract(filename, mime, content_b64) -> Document. Produces IF-4. Raises
UnsupportedFile -> 415.
"""

from __future__ import annotations

import base64
import binascii
import io
from collections.abc import Callable
from typing import Any

import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation

from app.detect.base import Document
from app.extract.docx_ooxml import extract_docx_ooxml_text
from app.extract.ocr import ocr_image_bytes, ocr_pdf_bytes
from app.extract.pptx_ooxml import extract_pptx_ooxml_text
from app.extract.xlsx import extract_xlsx_text

MAX_ATTACHMENT_BYTES = 5 * 1024 * 1024


class UnsupportedFile(Exception):
    """Raised for a filename/mime extract() cannot handle. Mapped to 415."""


def extract(filename: str, mime: str, content_b64: str) -> Document:
    raw = _decode_content(filename, content_b64)
    extractor = _extractor_for(filename, mime)
    text = extractor(raw)
    return Document(text=text, source_filename=filename, mime=mime)


def _decode_content(filename: str, content_b64: str) -> bytes:
    try:
        raw = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise UnsupportedFile(f"invalid base64 (filename={filename!r})") from exc

    if len(raw) > MAX_ATTACHMENT_BYTES:
        raise UnsupportedFile(f"file too large (filename={filename!r})")
    return raw


def _extractor_for(filename: str, mime: str) -> Callable[[bytes], str]:
    normalised_mime = mime.lower().split(";")[0].strip()
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if normalised_mime in {"text/plain", "text/csv"} or ext in {"txt", "csv"}:
        return _extract_text
    if normalised_mime in {"image/png", "image/jpeg"} or ext in {"png", "jpg", "jpeg"}:
        return _extract_image
    if normalised_mime == "application/pdf" or ext == "pdf":
        return _extract_pdf
    if (
        normalised_mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or ext == "docx"
    ):
        return _extract_docx
    if (
        normalised_mime
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        or ext == "pptx"
    ):
        return _extract_pptx
    if (
        normalised_mime
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or ext == "xlsx"
    ):
        return _extract_xlsx

    raise UnsupportedFile(f"{mime} (filename={filename!r})")


def _extract_text(raw: bytes) -> str:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedFile("text attachment is not valid UTF-8") from exc


def _extract_pdf(raw: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        return text if text.strip() else ocr_pdf_bytes(raw)
    except Exception as exc:
        raise UnsupportedFile("PDF text could not be extracted") from exc


def _extract_image(raw: bytes) -> str:
    try:
        return ocr_image_bytes(raw).text
    except Exception as exc:
        raise UnsupportedFile("image text could not be extracted") from exc


def _extract_docx(raw: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(raw))
    except Exception as exc:
        raise UnsupportedFile("DOCX text could not be extracted") from exc

    parts: list[str] = []
    for container in _iter_docx_story_containers(doc):
        parts.extend(p.text for p in container.paragraphs if p.text)
        for table in container.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append("\t".join(cells))
    ooxml_text = extract_docx_ooxml_text(raw)
    if ooxml_text:
        parts.append(ooxml_text)
    return "\n".join(parts)


def _iter_docx_story_containers(doc: Any) -> list[Any]:
    containers: list[Any] = [doc]
    seen = {id(doc._element)}
    for section in doc.sections:
        for attr in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            container = getattr(section, attr)
            marker = id(container._element)
            if marker in seen:
                continue
            seen.add(marker)
            containers.append(container)
    return containers


def _extract_pptx(raw: bytes) -> str:
    try:
        deck = Presentation(io.BytesIO(raw))
    except Exception as exc:
        raise UnsupportedFile("PPTX text could not be extracted") from exc

    parts: list[str] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            parts.append(slide.notes_slide.notes_text_frame.text)
    ooxml_text = extract_pptx_ooxml_text(raw)
    if ooxml_text:
        parts.append(ooxml_text)
    return "\n".join(parts)


def _extract_xlsx(raw: bytes) -> str:
    try:
        return extract_xlsx_text(raw)
    except Exception as exc:
        raise UnsupportedFile("XLSX text could not be extracted") from exc


__all__ = ["MAX_ATTACHMENT_BYTES", "UnsupportedFile", "extract"]
